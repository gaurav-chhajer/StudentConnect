import io
import unittest
from contextlib import closing
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from werkzeug.datastructures import FileStorage

from backend import resumes
from tests import test_web_flows


def pdf_file(text="Python SQL project", pages=1, encrypted=False):
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if text:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
                DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(("BT /F1 12 Tf 40 700 Td (" + text + ") Tj ET").encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("synthetic-test-password")
    result = io.BytesIO()
    writer.write(result)
    writer.close()
    return result.getvalue()


def docx_file():
    document = Document()
    document.add_paragraph("Python project")
    table = document.add_table(rows=1, cols=2)
    table.cell(0,0).text = "SQL"
    table.cell(0,1).text = "Internship"
    document.add_paragraph("Education final section")
    document.sections[0].header.paragraphs[0].text = "Candidate header"
    document.sections[0].footer.paragraphs[0].text = "Available immediately"
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


class ResumeDocumentTests(unittest.TestCase):
    def upload(self, content, filename):
        with closing(FileStorage(stream=io.BytesIO(content), filename=filename)) as upload:
            return resumes.read_upload(upload)

    def test_pdf_real_text_and_case_insensitive_suffix(self):
        content = pdf_file()
        text, original, suffix = self.upload(content, "resume.PDF")
        self.assertIn("Python SQL project", text)
        self.assertEqual(original, content)
        self.assertEqual(suffix, ".pdf")

    def test_docx_body_table_header_and_footer(self):
        text, _, suffix = self.upload(docx_file(), "resume.docx")
        for phrase in ("Python project","SQL","Internship","Education final section",
                       "Candidate header","Available immediately"):
            self.assertIn(phrase,text)
        self.assertLess(text.index("SQL"), text.index("Education final section"))
        self.assertEqual(suffix,".docx")

    def test_invalid_encrypted_scanned_and_oversized_files(self):
        cases = [(b"not PDF","resume.pdf"),(b"not ZIP","resume.docx"),
                 (pdf_file(encrypted=True),"resume.pdf"),(pdf_file(text=""),"resume.pdf"),
                 (pdf_file(pages=21),"resume.pdf"),(b"old word","resume.doc"),
                 (b"x"*(resumes.MAX_UPLOAD_BYTES+1),"resume.pdf")]
        for content,filename in cases:
            with self.subTest(filename=filename,size=len(content)), self.assertRaises(ValueError):
                self.upload(content,filename)

    def test_docx_expansion_limit_and_wrong_zip(self):
        with patch.object(resumes,"MAX_EXPANDED_BYTES",100),self.assertRaises(ValueError):
            self.upload(docx_file(),"resume.docx")
        archive = io.BytesIO()
        with ZipFile(archive,"w",ZIP_DEFLATED) as zip_file:
            zip_file.writestr("unrelated.txt","Not a Word document")
        with self.assertRaises(ValueError):
            self.upload(archive.getvalue(),"resume.docx")

    def test_pdf_stream_and_text_limits(self):
        with patch.object(resumes,"MAX_EXPANDED_BYTES",10),self.assertRaises(ValueError):
            self.upload(pdf_file(),"resume.pdf")
        with patch.object(resumes,"MAX_RESUME_CHARACTERS",10),self.assertRaises(ValueError):
            self.upload(pdf_file(),"resume.pdf")


class ResumeDocumentFlowTests(unittest.TestCase):
    setUp = test_web_flows.WebFlowTests.setUp
    seed_user = test_web_flows.WebFlowTests.seed_user
    post = test_web_flows.WebFlowTests.post

    def test_both_formats_save_private_originals_and_stale_scores(self):
        self.seed_user()
        for content,filename in ((pdf_file(),"../../resume.pdf"),(docx_file(),"resume.docx")):
            response = self.post("/resume",dict(resume_file=(io.BytesIO(content),filename),consent="on"))
            self.assertEqual(response.status_code,302)
            stored = self.conn.execute("SELECT resume_text FROM users WHERE id=1").fetchone()[0]
            self.assertIn("Python",stored)
            self.assertIn("SQL",stored)
        self.assertEqual(self.conn.execute("SELECT SUM(is_stale) FROM scores").fetchone()[0],2)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM resume_versions").fetchone()[0],2)
        for path in self.uploads.iterdir():
            self.assertIn(path.suffix,(".pdf",".docx"))
            self.assertEqual(self.client.get("/uploads/"+path.name).status_code,404)
        page = self.client.get("/resume")
        self.addCleanup(page.close)
        self.assertIn(b"Education final section",page.data)
        self.assertIn(b'.pdf,.docx,.txt',page.data)

    def test_bad_documents_do_not_overwrite_resume(self):
        self.seed_user()
        for filename in ("resume.pdf","resume.docx"):
            response = self.post("/resume",dict(resume_file=(io.BytesIO(b"bad file"),filename),consent="on"))
            self.assertEqual(response.status_code,400)
        self.assertEqual(self.conn.execute("SELECT resume_text FROM users WHERE id=1").fetchone()[0],"Python resume")
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM resume_versions").fetchone()[0],0)
        self.assertFalse(self.uploads.exists())
